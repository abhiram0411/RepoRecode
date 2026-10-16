"""Build the editable development report from verified repository evidence."""
from pathlib import Path
import json
import math
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'artifacts'
OUT.mkdir(exist_ok=True)
status_path = OUT / 'deployment.json'
status = json.loads(status_path.read_text()) if status_path.exists() else {}
live = status.get('status') == 'succeeded'
url = status.get('url', '') if live else ''
doc = Document()
# Remove decorative rules inherited from the bundled Word template.
for element in list(doc.styles.element.iter(qn('w:pBdr'))):
    element.getparent().remove(element)
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.8)
sec.footer_distance = Inches(.3)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption']:
    s = doc.styles[name]
    s.font.name = 'Calibri'
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.font.size = Pt(10.5)
    s.paragraph_format.space_after = Pt(7)
    s.paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(29)
doc.styles['Title'].font.bold = True
doc.styles['Heading 1'].font.size = Pt(20)
doc.styles['Heading 2'].font.size = Pt(13)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(10)
doc.styles['Caption'].font.size = Pt(9)
doc.styles['Caption'].font.italic = True
doc.core_properties.title = 'RepoRecode Development Report'
doc.core_properties.subject = 'Verified development release and remaining production scope'
doc.core_properties.author = 'RepoRecode'
doc.core_properties.keywords = 'RepoRecode, development, verification'
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('RepoRecode  |  16 October 2026  |  ').font.size = Pt(8)
f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'), 'PAGE'); footer._p.append(f)

def p(text, style=None):
    return doc.add_paragraph(text, style)

def h(text):
    doc.add_heading(text, 2)

def page(title):
    doc.add_page_break()
    doc.add_heading(title, 1)

def bullets(items):
    for text in items:
        p(text, 'List Bullet')

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, w in zip(t.columns, widths): col.width = Inches(w)
    for i, val in enumerate(headers): t.rows[0].cells[i].text = val
    for values in rows:
        cells = t.add_row().cells
        for i, val in enumerate(values): cells[i].text = str(val)
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el = OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'5'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
    pr.append(borders)
    for ri, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        nosplit = OxmlElement('w:cantSplit'); trpr.append(nosplit)
        if ri == 0:
            repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                el = OxmlElement('w:'+side); el.set(qn('w:w'),'100'); el.set(qn('w:type'),'dxa'); margins.append(el)
            cp.append(margins)
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E4EBEF' if ri == 0 else 'FFFFFF'); cp.append(shade)
            for par in cell.paragraphs:
                par.paragraph_format.space_after = Pt(2)
                par.paragraph_format.line_spacing = 1.0
                for run in par.runs:
                    run.font.size = Pt(9.5); run.bold = ri == 0
    p('')
    return t

def picture(name, width, caption):
    target = OUT / name
    if not target.exists(): return
    par = p(''); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run(); shape = run.add_picture(str(target), width=Inches(width))
    shape._inline.docPr.set('descr', caption)
    par.paragraph_format.keep_with_next = True
    p(caption, 'Caption')

def architecture():
    im = Image.new('RGB', (1500, 790), 'white'); d = ImageDraw.Draw(im)
    regular = ImageFont.truetype('C:/Windows/Fonts/calibri.ttf', 27)
    bold = ImageFont.truetype('C:/Windows/Fonts/calibrib.ttf', 29)
    boxes = [
      ((20,20,470,220), 'Browser workspace', ['React and Monaco', 'Lessons and progress', 'IndexedDB on this device']),
      ((530,20,980,220), 'Bounded workers', ['TypeScript source index', 'QuickJS evaluation', 'Sanitized React snapshot']),
      ((1040,20,1480,220), 'Imported source', ['Public GitHub or local files', 'Bounded immutable snapshot', 'No project install scripts']),
      ((20,450,470,680), 'Optional Supabase', ['Auth and owner protected rows', 'Projects and preferences', 'Encrypted GitHub grants']),
      ((530,450,980,680), 'Optional Node gateway', ['Verified bearer and origin', 'GitHub App and feedback', 'Quotas and administration']),
      ((1040,450,1480,680), 'External services', ['No paid model provider', 'GitHub repository API', 'Explicit alert destinations']),
    ]
    for box, title, lines in boxes:
        d.rounded_rectangle(box, radius=14, fill='#F4F7F9', outline='#6E818E', width=2)
        x,y,_,_=box; d.text((x+20,y+20),title,font=bold,fill='black')
        for i,line in enumerate(lines): d.text((x+20,y+68+i*38),line,font=regular,fill='#293E4A')
    def arrow(a,b,label=None):
        d.line([a,b],fill='#284B63',width=4)
        x,y=b; angle=math.atan2(b[1]-a[1],b[0]-a[0])
        d.polygon([(x,y),(x-14*math.cos(angle-.5),y-14*math.sin(angle-.5)),(x-14*math.cos(angle+.5),y-14*math.sin(angle+.5))],fill='#284B63')
    arrow((530,115),(470,115)); arrow((1040,115),(980,115))
    arrow((240,220),(750,450)); d.text((485,275),'Opt in',font=regular,fill='#293E4A')
    arrow((530,560),(470,560)); arrow((980,560),(1040,560))
    d.text((20,733),'Top row works locally. Bottom row requires owner configuration and live acceptance.',font=regular,fill='#293E4A')
    im.save(OUT/'architecture.png')

def evidence(name):
    path = OUT / name
    return json.loads(path.read_text(encoding='utf-8-sig')) if path.exists() else {}

units = evidence('unit-results.json')
browser = evidence('browser-results.json').get('stats', {})
audit = evidence('dependency-audit.json')
download = evidence('source-download.json')
if not units.get('success') or units.get('numFailedTests', 0) or browser.get('unexpected', 0):
    raise RuntimeError('Resolve failing checks before authoring a verified report.')
replacements = {
    '{{UNIT_TESTS}}': str(units.get('numPassedTests', 'unavailable')),
    '{{BROWSER_TESTS}}': str(browser.get('expected', 'unavailable')),
    '{{BROWSER_FAILURES}}': str(browser.get('unexpected', 'unavailable')),
    '{{PUBLIC_VERSION}}': str(status.get('version_number', 'unavailable')),
    '{{DOWNLOAD_STATUS}}': 'HTTP 200 and hash matched' if download.get('matchesLocal') and download.get('status') == 200 else 'Not verified',
    '{{VULNERABILITIES}}': str(audit.get('metadata', {}).get('vulnerabilities', {}).get('total', 'unavailable')),
}
content = (ROOT / 'docs' / 'ENGINEERING_REPORT.md').read_text(encoding='utf-8')
for key, value in replacements.items(): content = content.replace(key, value)
lines = iter(content.splitlines())
pending_rows = []
def flush_table():
    if pending_rows:
        table(pending_rows[0], pending_rows[1:], [2.0, 4.9])
        pending_rows.clear()
for line in lines:
    if line.startswith('|'):
        cells = [x.strip() for x in line.strip('|').split('|')]
        if not all(set(x) <= {'-', ' ', ':'} for x in cells): pending_rows.append(cells)
        continue
    flush_table()
    if not line.strip(): continue
    if line == '[architecture]':
        architecture(); picture('architecture.png', 6.9, 'Local learning and optional service boundaries')
    elif line.startswith('### '): h(line[4:])
    elif line.startswith('## '): page(line[3:])
    elif line.startswith('# '): doc.add_heading(line[2:], 0)
    elif line.startswith('- '): p(line[2:], 'List Bullet')
    else: p(line)
flush_table()
doc.save(OUT / 'RepoRecode-development-report.docx')
print('Created artifacts/RepoRecode-development-report.docx')
